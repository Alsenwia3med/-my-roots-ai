"""
Report-review point 21 as an editable workbook for ROOTS.

    python scripts/canonical/point21_workbook.py      (npm run evidence:point21-xlsx)

ROOTS review of 29 September 2026, item 1: the point 21 matrix is to be supplied in a form ROOTS
can work in and return, rather than a document to read.

Nothing here is authored. Every value is read from docs/m3/evidence/point21-matrix.json, which
`npm run evidence:point21` writes from the controlled content module and the real report builder,
so the workbook and the Markdown matrix cannot say different things.

The four decision columns are empty and unlocked. Every column that carries controlled content is
locked and greyed, and the sheet is protected without a password: the protection is there to stop
an accidental edit to approved wording, not to stop ROOTS doing anything they mean to do
(Review > Unprotect Sheet, no password).

Output: docs/m3/evidence/Point21_Content_Matrix_FOR_ROOTS.xlsx

Requires: pip install openpyxl
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "docs" / "m3" / "evidence" / "point21-matrix.json"
OUT = ROOT / "docs" / "m3" / "evidence" / "Point21_Content_Matrix_FOR_ROOTS.xlsx"

NAVY = "1A2A4A"
GREY = "F1F3F6"
CREAM = "FFF8E7"
LINE = "D9E0E8"

HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY = Font(size=10)
WRAP = Alignment(vertical="top", wrap_text=True)
THIN = Border()

# (heading, width, editable)
COLUMNS = [
    ("#", 5, False),
    ("Domain", 8, False),
    ("Domain label (C-03 §3)", 26, False),
    ("Classification", 15, False),
    ("C-02 range", 11, False),
    ("Approved domain meaning (C-03 §3)", 46, False),
    ("Approved classification explanation (C-03 §3)", 46, False),
    ("Exact current rendered text", 52, False),
    ("ROOTS: additional content required?", 18, True),
    ("ROOTS: approved domain-specific interpretation", 52, True),
    ("ROOTS: content identifier", 20, True),
    ("ROOTS: notes", 34, True),
]

READ_ME = [
    ("ROOTS-AI™ — report-review point 21 content matrix", True),
    ("", False),
    ("What this is", True),
    (
        "The seven-domain x four-classification matrix the ROOTS decision of 29 September 2026 asks for, "
        "in editable form (review of 29 September 2026, item 1).",
        False,
    ),
    ("", False),
    ("What ROOTS is asked to determine", True),
    (
        "Per the decision: \"ROOTS will determine which combinations genuinely require additional approved "
        "content.\" The four columns headed ROOTS: are empty and unlocked. Everything else is locked, because "
        "it is controlled content reproduced verbatim.",
        False,
    ),
    ("", False),
    (
        "Column I takes Yes or No from a drop-down. Where it is Yes, column J takes the approved wording and "
        "column K its content identifier. Column L is for anything else.",
        False,
    ),
    ("", False),
    ("Where the values come from", True),
    (
        "Nothing in this workbook was authored by the vendor. Domain labels, domain meanings and classification "
        "explanations are read from the controlled C-03 content module; the C-02 ranges come from the "
        "Classifications sheet; and the rendered text in column H is produced by the real report builder, so it "
        "is exactly what a participant sees today.",
        False,
    ),
    ("", False),
    (
        "All 28 combinations currently have no domain-specific interpretation in C-03 v1.0.1. The controlled pack "
        "supplies per-domain meanings and four per-classification explanations, and nothing that combines the two. "
        "That pairing is the mechanism behind the repetition point 21 describes: any two domains sharing a "
        "classification necessarily read identically.",
        False,
    ),
    ("", False),
    ("Constraints the decision records as remaining in force", True),
    ("Any addition must preserve:", False),
    (
        "• the four existing classification explanations and their medical boundaries, unless a controlled "
        "amendment expressly changes them;",
        False,
    ),
    (
        "• deterministic classifications — no added content may alter how a score or classification is produced;",
        False,
    ),
    ("• non-diagnostic framing and the approved safety guidance.", False),
    ("", False),
    (
        "Review point 21 also excludes solving the repetition \"by allowing unrestricted AI rewriting\", and the "
        "decision forbids newly authored physiological or medical claims. Neither route is open to the vendor, "
        "which is why the interpretation column is empty rather than filled.",
        False,
    ),
    ("", False),
    ("Regenerating this file", True),
    (
        "npm run evidence:point21 && npm run evidence:point21-xlsx — which overwrites it. Return a completed copy "
        "under a different name so nothing ROOTS has entered can be lost to a regeneration.",
        False,
    ),
]


def read_me_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Read me", 0)
    ws.column_dimensions["A"].width = 118
    ws.sheet_view.showGridLines = False
    for i, (text, heading) in enumerate(READ_ME, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(bold=heading, size=12 if i == 1 else 10, color=NAVY if heading else "000000")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = None if len(text) < 100 else 15 * (len(text) // 100 + 1)


def matrix_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Matrix")
    ws.freeze_panes = "A2"

    for i, (heading, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=heading)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34

    for r, row in enumerate(data["rows"], start=2):
        values = [
            row["n"],
            row["domain_id"],
            row["domain_label"],
            row["classification"],
            row["range"],
            row["domain_meaning"],
            row["classification_explanation"],
            row["rendered_text"] or "not produced by any Golden Test",
            None,
            None,
            None,
            None,
        ]
        for i, (value, (_, _, editable)) in enumerate(zip(values, COLUMNS), start=1):
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = BODY
            cell.alignment = WRAP
            cell.protection = Protection(locked=not editable)
            cell.fill = PatternFill("solid", fgColor=CREAM if editable else GREY)
        ws.row_dimensions[r].height = 58

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showDropDown=False)
    yes_no.error = "Enter Yes or No."
    yes_no.prompt = "Does this combination require additional approved content?"
    ws.add_data_validation(yes_no)
    yes_no.add(f"I2:I{len(data['rows']) + 1}")

    # No password: this guards against a slip, not against ROOTS.
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def reference_sheet(wb: Workbook, data: dict) -> None:
    ws = wb.create_sheet("Approved content")
    ws.sheet_view.showGridLines = False

    def header(row: int, cells: list[str], widths: list[int]) -> None:
        for i, (text, width) in enumerate(zip(cells, widths), start=1):
            c = ws.cell(row=row, column=i, value=text)
            c.font = HEADER_FONT
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width

    ws.cell(row=1, column=1, value="The four approved classification explanations (C-03 §3)").font = Font(
        bold=True, color=NAVY, size=11
    )
    header(2, ["Classification", "C-02 range", "C-02 approved interpretation", "C-03 participant-facing explanation"], [16, 26, 48, 68])
    r = 3
    for band in data["bands"]:
        for i, value in enumerate(
            [band["label"], band["range"], band["interpretation"], data["classification_explanations"][band["label"]]],
            start=1,
        ):
            c = ws.cell(row=r, column=i, value=value)
            c.font = BODY
            c.alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="The seven approved domain meanings (C-03 §3)").font = Font(bold=True, color=NAVY, size=11)
    r += 1
    header(r, ["Domain", "Label", "Approved meaning"], [16, 26, 68])
    r += 1
    for d in data["domains"]:
        for i, value in enumerate([d["id"], d["label"], d["meaning"]], start=1):
            c = ws.cell(row=r, column=i, value=value)
            c.font = BODY
            c.alignment = WRAP
        ws.row_dimensions[r].height = 30
        r += 1

    ws.protection.sheet = True


def main() -> None:
    if not DATA.exists():
        raise SystemExit(f"{DATA.relative_to(ROOT)} not found — run `npm run evidence:point21` first")

    data = json.loads(DATA.read_text(encoding="utf-8"))

    wb = Workbook()
    wb.remove(wb.active)
    read_me_sheet(wb)
    matrix_sheet(wb, data)
    reference_sheet(wb, data)
    wb.active = 0

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    filled = sum(1 for row in data["rows"] if row["rendered_text"])
    print(f"wrote {OUT.relative_to(ROOT)}")
    print(f"{len(data['rows'])} combinations; {filled} with a rendered sample; 4 decision columns left empty")


if __name__ == "__main__":
    main()
