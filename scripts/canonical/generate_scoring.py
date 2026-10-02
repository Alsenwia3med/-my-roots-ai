#!/usr/bin/env python3
"""
Generate canonical scoring data from C-02 Excel workbook.
This script reads the controlled C-02 workbook and generates JSON files
for the scoring engine to use.
"""

import json
import hashlib
import sys
from pathlib import Path
from typing import Dict, List, Any
import openpyxl

def calculate_sha256(file_path: str) -> str:
    """Calculate SHA-256 hash of a file."""
    sha256_hash = hashlib.sha256()
    with open(file_path, "rb") as f:
        for byte_block in iter(lambda: f.read(4096), b""):
            sha256_hash.update(byte_block)
    return sha256_hash.hexdigest()

def read_workbook(file_path: str) -> openpyxl.Workbook:
    """Read the Excel workbook."""
    return openpyxl.load_workbook(file_path, data_only=True)

def find_header_row(sheet, header_name: str) -> int:
    """Find the row containing a specific header."""
    for idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if row and header_name in [str(cell) if cell else '' for cell in row]:
            return idx
    raise ValueError(f"Header '{header_name}' not found")

def extract_domains(sheet) -> List[Dict[str, Any]]:
    """Extract domain information from the Domains sheet."""
    domains = []
    header_row = find_header_row(sheet, 'domain_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        domain = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                domain[headers[idx]] = value

        if domain:
            domains.append(domain)

    return domains

def extract_question_mapping(sheet) -> List[Dict[str, Any]]:
    """Extract question to domain mapping from the Question_Mapping sheet."""
    mapping = []
    header_row = find_header_row(sheet, 'question_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        item = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                item[headers[idx]] = value

        if item:
            mapping.append(item)

    return mapping

def extract_option_points(sheet) -> Dict[str, List[Dict[str, Any]]]:
    """Extract option point values from the Option_Points sheet."""
    option_points = {}
    header_row = find_header_row(sheet, 'option_set_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    current_set_id = None
    current_options = []

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            if current_set_id and current_options:
                option_points[current_set_id] = current_options
                current_set_id = None
                current_options = []
            continue

        option = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                option[headers[idx]] = value

        if option.get('option_set_id') != current_set_id:
            if current_set_id and current_options:
                option_points[current_set_id] = current_options
            current_set_id = option.get('option_set_id')
            current_options = []

        if option:
            current_options.append(option)

    if current_set_id and current_options:
        option_points[current_set_id] = current_options

    return option_points

def extract_formulas(sheet) -> List[Dict[str, Any]]:
    """Extract scoring formulas from the Formulas sheet."""
    formulas = []
    header_row = find_header_row(sheet, 'rule_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        formula = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                formula[headers[idx]] = value

        if formula:
            formulas.append(formula)

    return formulas

def extract_protective_factors(sheet) -> List[Dict[str, Any]]:
    """Extract protective factors from the Protective_Factors sheet."""
    factors = []
    header_row = find_header_row(sheet, 'factor_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        factor = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                factor[headers[idx]] = value

        if factor:
            factors.append(factor)

    return factors

def extract_classifications(sheet) -> List[Dict[str, Any]]:
    """Extract classification bands from the Classifications sheet."""
    classifications = []
    header_row = find_header_row(sheet, 'scale')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        classification = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                classification[headers[idx]] = value

        if classification:
            classifications.append(classification)

    return classifications

def extract_golden_tests(sheet) -> List[Dict[str, Any]]:
    """Extract golden test cases from the Golden_Tests sheet."""
    tests = []
    header_row = find_header_row(sheet, 'test_id')

    headers = list(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))[0]
    headers = [h.strip().lower().replace(' ', '_') if h else '' for h in headers]

    for row in sheet.iter_rows(min_row=header_row + 2, values_only=True):
        if not row[0]:
            continue

        test = {}
        for idx, value in enumerate(row):
            if idx < len(headers) and headers[idx]:
                # Try to parse JSON strings
                if isinstance(value, str) and value.startswith('{'):
                    try:
                        test[headers[idx]] = json.loads(value)
                    except:
                        test[headers[idx]] = value
                else:
                    test[headers[idx]] = value

        if test:
            tests.append(test)

    return tests

def main():
    workbook_path = r"c:\Users\user\Downloads\03_ROOTS_AI_C02_Canonical_Scoring_Rules_and_Golden_Tests_v1.0.1_CORRECTED.xlsx"
    output_dir = Path("C:/Users/user/Desktop/works/2026/RootsAi/my-roots-ai/src/lib/scoring")
    output_dir.mkdir(parents=True, exist_ok=True)

    # Calculate SHA-256
    sha256 = calculate_sha256(workbook_path)
    print(f"Workbook SHA-256: {sha256}")

    # Read workbook
    wb = read_workbook(workbook_path)

    # Extract data
    domains = extract_domains(wb['Domains'])
    question_mapping = extract_question_mapping(wb['Question_Mapping'])
    option_points = extract_option_points(wb['Option_Points'])
    formulas = extract_formulas(wb['Formulas'])
    protective_factors = extract_protective_factors(wb['Protective_Factors'])
    classifications = extract_classifications(wb['Classifications'])
    golden_tests = extract_golden_tests(wb['Golden_Tests'])

    # Build complete scoring ruleset
    scoring_ruleset = {
        "metadata": {
            "document": "C-02 v1.0.1 CORRECTED",
            "scoring_version": "1.0.1",
            "sha256": sha256,
            "generated_at": "2026-10-02"
        },
        "domains": domains,
        "question_mapping": question_mapping,
        "option_points": option_points,
        "formulas": formulas,
        "protective_factors": protective_factors,
        "classifications": classifications,
        "golden_tests": golden_tests
    }

    # Write output
    output_file = output_dir / "c02-ruleset.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(scoring_ruleset, f, indent=2, ensure_ascii=False)

    print(f"Generated {output_file}")
    print(f"Domains: {len(domains)}")
    print(f"Question Mappings: {len(question_mapping)}")
    print(f"Option Point Sets: {len(option_points)}")
    print(f"Formulas: {len(formulas)}")
    print(f"Protective Factors: {len(protective_factors)}")
    print(f"Classifications: {len(classifications)}")
    print(f"Golden Tests: {len(golden_tests)}")

if __name__ == "__main__":
    main()
