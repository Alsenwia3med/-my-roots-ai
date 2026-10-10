"""
Content fingerprint of the controlled XLSX workbooks — M2 closure, provenance item 1.

A file-level SHA-256 changes whenever a workbook is re-saved or re-packaged (zip compression,
timestamps, document properties, style tables), even when not a single cell changes. This
fingerprints only the canonical CONTENT: for every sheet, every non-empty cell's address and
either its value or, for a formula cell, its formula text. A formula's cached result is not
content — it depends on which program last saved the file (a re-save by a tool that does not
recalculate drops it) — so it is excluded. Formatting, styles, column widths and file metadata
are ignored.

    # fingerprint one or more workbooks
    python scripts/canonical/content_fingerprint.py FILE.xlsx [FILE2.xlsx ...]

    # compare two copies of the same workbook, cell by cell
    python scripts/canonical/content_fingerprint.py --compare A.xlsx B.xlsx

Two copies with the same content fingerprint contain identical canonical content. If they
differ, --compare lists every differing sheet and cell (address, value in A, value in B).

Requires: pip install openpyxl
"""

import hashlib
import json
import sys
from datetime import date, datetime

import openpyxl


def norm(v):
    """A cell value as canonical text, independent of how the file was written."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def read_cells(path):
    """{sheet: {cell: {"value": ..., "formula": ...}}} for every non-empty cell, sheets in order."""
    values = openpyxl.load_workbook(path, data_only=True)  # stored (cached) values
    formulas = openpyxl.load_workbook(path, data_only=False)  # formula text where present
    out = {}
    for ws in values.worksheets:
        fws = formulas[ws.title]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = norm(c.value)
                f = fws[c.coordinate].value
                f = f if isinstance(f, str) and f.startswith("=") else None
                if v is None and f is None:
                    continue
                cells[c.coordinate] = {"formula": f} if f else {"value": v}
        out[ws.title] = cells
    return out


def digest(obj):
    return hashlib.sha256(json.dumps(obj, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()


def fingerprint(path):
    cells = read_cells(path)
    sheets = {name: {"cells": len(c), "content_sha256": digest(c)} for name, c in cells.items()}
    return {
        "file": path,
        "file_sha256": hashlib.sha256(open(path, "rb").read()).hexdigest(),
        "sheet_order": list(cells),
        "sheets": sheets,
        "content_sha256": digest({"order": list(cells), "sheets": cells}),
    }


def compare(a, b):
    ca, cb = read_cells(a), read_cells(b)
    diffs = []
    if list(ca) != list(cb):
        diffs.append(f"sheet order differs: {list(ca)} vs {list(cb)}")
    for name in dict.fromkeys(list(ca) + list(cb)):
        sa, sb = ca.get(name), cb.get(name)
        if sa is None or sb is None:
            diffs.append(f"sheet '{name}' only in {'B' if sa is None else 'A'}")
            continue
        for addr in sorted(set(sa) | set(sb)):
            if sa.get(addr) != sb.get(addr):
                diffs.append(f"{name}!{addr}: A={sa.get(addr)} | B={sb.get(addr)}")
    return diffs


def main():
    args = sys.argv[1:]
    if args[:1] == ["--compare"] and len(args) == 3:
        diffs = compare(args[1], args[2])
        fa, fb = fingerprint(args[1]), fingerprint(args[2])
        print(f"A file SHA-256   {fa['file_sha256']}\nB file SHA-256   {fb['file_sha256']}")
        print(f"A content SHA-256 {fa['content_sha256']}\nB content SHA-256 {fb['content_sha256']}")
        if diffs:
            print(f"CONTENT DIFFERS — {len(diffs)} difference(s):")
            for d in diffs:
                print("  " + d)
            sys.exit(1)
        print("CONTENT IDENTICAL — every sheet, cell value and formula matches.")
        return
    if not args:
        sys.exit(__doc__)
    print(json.dumps([fingerprint(p) for p in args], indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
